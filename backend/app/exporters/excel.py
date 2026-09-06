from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MAX_CELL = 32000
MAX_ROWS = 1_000_000


def safe_filename(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", name).strip(" ._")
    return cleaned[:80] or "project"


def sanitize_cell(value: Any, *, numeric: bool = False) -> Any:
    if value is None:
        return None if numeric else ""
    if numeric:
        if value == "" or value == "unknown":
            return None
        return value
    text = str(value)
    if len(text) > MAX_CELL:
        text = text[: MAX_CELL - 20] + "\n[truncated; see JSON]"
    if text.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def _header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for idx, _ in enumerate(headers, start=1):
        cell = ws.cell(1, idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(idx)].width = min(28, max(14, len(_) + 2))


def _row(ws: Worksheet, values: list[Any], numeric_idx: set[int] | None = None) -> None:
    numeric_idx = numeric_idx or set()
    ws.append([sanitize_cell(v, numeric=i in numeric_idx) for i, v in enumerate(values)])
    for cell in ws[ws.max_row]:
        cell.alignment = WRAP


def write_workbook(path: Path, payload: dict[str, Any]) -> None:
    wb = Workbook()
    _readme(wb.active, payload)
    _candidates(wb.create_sheet("Candidates"), payload)
    _subset(wb.create_sheet("Recommended"), payload, "recommended")
    _subset(wb.create_sheet("Needs_review"), payload, "needs_review")
    _subset(wb.create_sheet("Excluded"), payload, "excluded")
    _evidence(wb.create_sheet("Evidence"), payload)
    _samples(wb.create_sheet("Samples"), payload)
    _queries(wb.create_sheet("Queries"), payload)
    _run_info(wb.create_sheet("Run_info"), payload)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _readme(ws: Worksheet, payload: dict[str, Any]) -> None:
    ws.title = "README"
    run = payload["run"]
    spec = payload.get("spec") or {}
    lines = [
        ("软件", "GEOScout"),
        ("课题", spec.get("original_request") or payload.get("project_name")),
        ("运行时间", run.get("created_at")),
        ("任务状态", run.get("status")),
        ("停止原因", run.get("stop_reason") or ""),
        ("完整性", "完整" if run.get("status") == "completed" else "部分完成，见 Run_info 与 Queries"),
        ("推荐", "全部硬条件 pass，且完成深度核验与复核"),
        ("待核实", "硬条件 unknown、冲突或核验未完成"),
        ("排除", "至少一条硬条件有已核验的明确 fail"),
        ("限制", "不声称穷尽 GEO；未知计数留空而不是 0；模型不得编造 accession"),
        ("演示模式", "是" if payload.get("demo") else "否（真实检索失败不会改用演示数据）"),
    ]
    _header(ws, ["项", "内容"])
    for row in lines:
        _row(ws, list(row))
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 80


def _candidate_headers() -> list[str]:
    return [
        "GSE",
        "标题",
        "官方链接",
        "PMID",
        "物种",
        "组织",
        "疾病",
        "技术",
        "平台",
        "GSM数",
        "已确认供体数",
        "分组与组内供体",
        "细胞数",
        "对照类型",
        "临床字段",
        "处理后数据",
        "原始数据",
        "最终类别",
        "硬条件未知数",
        "软条件满足度",
        "核验完成状态",
        "理由",
        "疑点",
        "相关系列",
        "人工覆盖",
        "采用GSM子集",
        "未知硬条件",
    ]


def _candidates(ws: Worksheet, payload: dict[str, Any]) -> None:
    _header(ws, _candidate_headers())
    numeric = {9, 10, 12, 18, 19}
    for row in payload.get("candidates") or []:
        _row(ws, [_candidate_values(row)[i] for i in range(len(_candidate_headers()))], numeric)


def _subset(ws: Worksheet, payload: dict[str, Any], category: str) -> None:
    _header(ws, _candidate_headers())
    numeric = {9, 10, 12, 18, 19}
    for row in payload.get("candidates") or []:
        if row.get("category") == category:
            _row(ws, _candidate_values(row), numeric)


def _candidate_values(row: dict[str, Any]) -> list[Any]:
    return [
        row.get("gse"),
        row.get("title"),
        row.get("url"),
        ",".join(row.get("pubmed_ids") or []),
        row.get("taxon"),
        row.get("tissue"),
        row.get("disease"),
        row.get("gdstype"),
        row.get("gpl"),
        row.get("gsm_count"),
        row.get("independent_donors"),
        json.dumps(row.get("donors_per_group") or {}, ensure_ascii=False),
        row.get("cell_count"),
        row.get("control_type"),
        ",".join(row.get("clinical_fields") or []),
        row.get("processed_data"),
        row.get("raw_data"),
        row.get("category"),
        row.get("hard_unknowns"),
        row.get("soft_score"),
        row.get("verification_status"),
        row.get("reason"),
        row.get("concerns"),
        row.get("relations"),
        "yes" if row.get("overridden") else "no",
        ",".join(row.get("applicable_gsms") or []),
        ",".join(row.get("unknown_hard") or []),
    ]


def _evidence(ws: Worksheet, payload: dict[str, Any]) -> None:
    _header(ws, ["GSE", "条件", "阶段", "判断", "理由", "引文", "evidence_id", "来源URL", "字段路径", "判断者", "支持文本", "合格GSM"])
    extra_json: list[dict[str, Any]] = []
    rows = list(payload.get("assessments") or [])
    rows.sort(key=lambda r: (0 if r.get("stage") == "final" else 1, str(r.get("gse")), str(r.get("criterion_id"))))
    for row in rows:
        quote = row.get("quote") or ""
        if len(quote) > 3000:
            extra_json.append(row)
            quote = quote[:3000] + " [see sidecar JSON]"
        _row(
            ws,
            [
                row.get("gse"),
                row.get("criterion_id"),
                row.get("stage"),
                row.get("verdict"),
                row.get("reason"),
                quote,
                ",".join(row.get("evidence_ids") or []),
                row.get("source_url"),
                row.get("field_path"),
                row.get("judge_source") or row.get("actor") or "",
                row.get("support_text") or "",
                ",".join(row.get("qualifying_gsms") or []),
            ],
        )
    payload["_long_evidence"] = extra_json


def _samples(ws: Worksheet, payload: dict[str, Any]) -> None:
    _header(ws, ["GSE", "GSM", "标题", "物种", "来源", "供体键", "特征", "覆盖不完整"])
    for row in payload.get("samples") or []:
        _row(
            ws,
            [
                row.get("gse"),
                row.get("gsm"),
                row.get("title"),
                row.get("organism"),
                row.get("source_name"),
                row.get("donor_key"),
                json.dumps(row.get("characteristics") or [], ensure_ascii=False)[:5000],
                "yes" if row.get("coverage_incomplete") else "no",
            ],
        )


def _queries(ws: Worksheet, payload: dict[str, Any]) -> None:
    _header(ws, ["轮次", "来源", "检索式", "NCBI翻译", "命中数", "新增唯一GSE", "状态", "截断", "错误", "时间"])
    numeric = {0, 4, 5}
    for row in payload.get("queries") or []:
        _row(
            ws,
            [
                row.get("round_no"),
                row.get("source"),
                row.get("term"),
                row.get("query_translation"),
                row.get("hit_count"),
                row.get("new_unique_gse"),
                row.get("status"),
                "yes" if row.get("truncated") else "no",
                row.get("error_message"),
                row.get("finished_at"),
            ],
            numeric,
        )


def _run_info(ws: Worksheet, payload: dict[str, Any]) -> None:
    _header(ws, ["键", "值"])
    info = payload.get("run_info") or {}
    for key, value in info.items():
        _row(ws, [key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])


def read_candidate_gses(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Candidates"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = headers.index("GSE")
    values = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx]:
            values.append(str(row[idx]))
    wb.close()
    return values


def read_sheet_maps(path: Path, sheet: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows)]
    out = []
    for row in rows:
        out.append({headers[i]: row[i] for i in range(len(headers))})
    wb.close()
    return out
